"""
ingestion.py — parse, chunk, embed, and store documents into ChromaDB
Supports: PDF, DOCX, PPTX, TXT
"""
import os, re
import fitz                            # PyMuPDF
import docx as _docx
from pptx import Presentation
from sentence_transformers import SentenceTransformer
import chromadb

EMBED_MODEL   = "BAAI/bge-base-en-v1.5"
CHUNK_SIZE    = 400    # approx words per chunk
CHUNK_OVERLAP = 80     # words to carry over for context continuity

embedder   = SentenceTransformer(EMBED_MODEL)
_db_client = chromadb.PersistentClient(path="./chroma_db")
collection = _db_client.get_or_create_collection(
    name="study_docs",
    metadata={"hnsw:space": "cosine"},
)

# ── Parsers ───────────────────────────────────────────────────────────────────

def _parse_pdf(path: str, display_name: str) -> list[dict]:
    doc   = fitz.open(path)
    pages = []
    for i, page in enumerate(doc):
        text = page.get_text("text").strip()
        if text:
            pages.append({"text": text, "page": i + 1, "source": display_name})
    doc.close()
    return pages

def _parse_docx(path: str, display_name: str) -> list[dict]:
    doc        = _docx.Document(path)
    paras      = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    group_size = 20
    pages      = []
    for i in range(0, len(paras), group_size):
        text = " ".join(paras[i : i + group_size])
        pages.append({"text": text, "page": i // group_size + 1, "source": display_name})
    return pages

def _parse_pptx(path: str, display_name: str) -> list[dict]:
    prs   = Presentation(path)
    pages = []
    for i, slide in enumerate(prs.slides):
        texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        if texts:
            pages.append({"text": " ".join(texts), "page": i + 1, "source": display_name})
    return pages

def _parse_txt(path: str, display_name: str) -> list[dict]:
    with open(path, encoding="utf-8", errors="replace") as f:
        lines = f.readlines()
    group_size = 50
    pages      = []
    for i in range(0, len(lines), group_size):
        text = "".join(lines[i : i + group_size]).strip()
        if text:
            pages.append({"text": text, "page": i // group_size + 1, "source": display_name})
    return pages

def _parse_file(path: str, display_name: str) -> list[dict]:
    ext = os.path.splitext(path)[1].lower()
    parsers = {
        ".pdf":  _parse_pdf,
        ".docx": _parse_docx,
        ".pptx": _parse_pptx,
        ".txt":  _parse_txt,
    }
    if ext not in parsers:
        raise ValueError(f"Unsupported file type: {ext}")
    return parsers[ext](path, display_name)

# ── Semantic chunker ──────────────────────────────────────────────────────────

def _split_sentences(text: str) -> list[str]:
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]

def _semantic_chunk(pages: list[dict]) -> list[dict]:
    """
    Split each page's text into overlapping word-budget chunks.
    Each chunk carries: text, page, source, chunk_index.
    """
    chunks = []
    for page in pages:
        sentences  = _split_sentences(page["text"])
        word_buf   = []
        chunk_idx  = 0

        for sent in sentences:
            sent_words = sent.split()
            if len(word_buf) + len(sent_words) > CHUNK_SIZE and word_buf:
                chunks.append({
                    "text":        " ".join(word_buf),
                    "page":        page["page"],
                    "source":      page["source"],
                    "chunk_index": chunk_idx,
                })
                chunk_idx += 1
                word_buf   = word_buf[-CHUNK_OVERLAP:] + sent_words
            else:
                word_buf.extend(sent_words)

        if word_buf:
            chunks.append({
                "text":        " ".join(word_buf),
                "page":        page["page"],
                "source":      page["source"],
                "chunk_index": chunk_idx,
            })
    return chunks

# ── Public API ────────────────────────────────────────────────────────────────

def ingest_file(path: str, original_name: str | None = None) -> int:
    """
    Parse, chunk, embed, and upsert a document into ChromaDB.
    Returns the number of chunks stored.

    Args:
        path:          Filesystem path to the temp file.
        original_name: Display name to store as metadata (defaults to basename).
    """
    display_name = original_name or os.path.basename(path)

    # Remove any previously stored chunks for this document
    try:
        existing = collection.get(where={"source": display_name})
        if existing["ids"]:
            collection.delete(ids=existing["ids"])
    except Exception:
        pass

    pages  = _parse_file(path, display_name)
    chunks = _semantic_chunk(pages)

    if not chunks:
        return 0

    texts      = [c["text"] for c in chunks]
    embeddings = embedder.encode(texts, show_progress_bar=False).tolist()
    ids        = [f"{display_name}__chunk_{i}" for i in range(len(chunks))]
    metadatas  = [{"source": c["source"], "page": c["page"]} for c in chunks]

    collection.upsert(
        ids=ids,
        documents=texts,
        embeddings=embeddings,
        metadatas=metadatas,
    )

    print(f"[ingestion] {display_name} → {len(chunks)} chunks stored")
    return len(chunks)


if __name__ == "__main__":
    import sys
    for f in sys.argv[1:]:
        n = ingest_file(f)
        print(f"Ingested {f}: {n} chunks")
